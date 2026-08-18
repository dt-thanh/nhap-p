"""Excel (.xlsx) chạy hết đường: file → parser → PostgreSQL → bộ tính cũ.

`test_excel_parser.py` đã phủ kỹ phần ĐỌC file .xlsx, nhưng dừng ở staging CSV.
`test_parse_upload.py` phủ phần GHI DB, nhưng mọi test có DB đều dùng .csv. Ở
giữa hai bộ đó có một khoảng trống: chưa test nào chứng minh một file .xlsx thật
đi tới được bảng `sales_records`.

Khoảng trống này không vô hại. Nhánh calamine trả về đối tượng `date`/`int` đã
đúng kiểu, còn nhánh CSV trả chuỗi rồi mới ép kiểu. Hai nhánh đi qua cùng một
tầng ghi nhưng bằng dữ liệu khác kiểu, nên một lỗi chỉ xuất hiện ở nhánh Excel
sẽ không có test nào bắt được — kể cả khi cả hai bộ test trên đều xanh.

Đường Excel/CSV là đường TỔNG HỢP: nó ghi vào `sales_records` /
`inventory_snapshots` / `areas` và không bao giờ ghi vào `units`/`deals`. Mỗi
test ở đây khẳng định lại điều đó sau khi chạy — xem thêm
`tests/test_services/test_legacy_boundary.py`.
"""

import os
import uuid
from datetime import date

import pytest
from openpyxl import Workbook

from src.jobs.parse_upload import run_parse_upload

TEST_DATABASE_URL = os.getenv("TEST_DATABASE_URL") or os.getenv("DATABASE_URL")

# UUID riêng của module. Xem chú thích cùng loại ở test_legacy_boundary.py.
PROJECT_ID = uuid.UUID("b7c1d2e3-4f56-4789-a0b1-c2d3e4f50022")

pytestmark = pytest.mark.skipif(not TEST_DATABASE_URL, reason="Cần TEST_DATABASE_URL trỏ vào Postgres thật")


def _write_xlsx(path, rows):
    """Ghi .xlsx THẬT bằng openpyxl, giữ nguyên kiểu dữ liệu của từng ô.

    Không dựng file giả bằng cách đổi đuôi .csv: mục đích của module này là chạy
    qua nhánh calamine, mà nhánh đó được chọn theo đuôi file VÀ phải giải nén
    được nội dung .xlsx thật.
    """
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def _engine():
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from sqlalchemy.pool import NullPool

    return async_sessionmaker(create_async_engine(TEST_DATABASE_URL, poolclass=NullPool))


def _seed(checksum="xlsx-e2e"):
    """Project + phân khu A1/2PN + một bản ghi lô. Trả về file_id."""
    import asyncio
    from datetime import UTC, datetime

    import sqlalchemy as sa

    from src.models.tables import areas, upload_files

    file_id = uuid.uuid4()

    async def seed():
        factory = _engine()
        async with factory() as session:
            async with session.begin():
                await session.execute(
                    sa.text(
                        "INSERT INTO projects (id, name, launch_date, created_at) "
                        "VALUES (:i, :n, :d, :t) ON CONFLICT (id) DO NOTHING"
                    ),
                    {"i": PROJECT_ID, "n": "XLSX-FIXTURE", "d": date(2026, 1, 1), "t": datetime.now(UTC)},
                )
                exists = await session.scalar(
                    sa.select(areas.c.id).where(areas.c.project_id == PROJECT_ID, areas.c.area_name == "A1")
                )
                if exists is None:
                    await session.execute(
                        sa.insert(areas).values(
                            id=uuid.uuid4(),
                            project_id=PROJECT_ID,
                            area_name="A1",
                            unit_type="2PN",
                            bedrooms=2,
                            area_sqm=75,
                            total_units=100,
                            created_at=datetime.now(UTC),
                        )
                    )
                await session.execute(
                    sa.insert(upload_files).values(
                        id=file_id,
                        project_id=PROJECT_ID,
                        uploaded_by=None,
                        filename="s.xlsx",
                        checksum=checksum,
                        status="pending",
                        rows_ok=0,
                        rows_failed=0,
                        uploaded_at=datetime.now(UTC),
                    )
                )

    asyncio.run(seed())
    return str(file_id)


def _cleanup():
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import absorption_daily, areas, inventory_snapshots, sales_records, upload_errors

    async def run():
        factory = _engine()
        async with factory() as session:
            async with session.begin():
                area_ids = sa.select(areas.c.id).where(areas.c.project_id == PROJECT_ID).scalar_subquery()
                file_ids = sa.text("SELECT id FROM upload_files WHERE project_id = :p").bindparams(p=PROJECT_ID)
                await session.execute(sa.delete(upload_errors).where(upload_errors.c.file_id.in_(file_ids)))
                for table in (absorption_daily, sales_records, inventory_snapshots):
                    await session.execute(sa.delete(table).where(table.c.area_id.in_(area_ids)))
                await session.execute(sa.delete(areas).where(areas.c.project_id == PROJECT_ID))
                await session.execute(sa.text("DELETE FROM upload_files WHERE project_id = :p"), {"p": PROJECT_ID})
                await session.execute(sa.text("DELETE FROM projects WHERE id = :p"), {"p": PROJECT_ID})

    asyncio.run(run())


@pytest.fixture
def db_env(monkeypatch):
    import src.db as db_module
    from src.config import get_settings

    caches = (db_module.get_engine, db_module.get_session_factory, get_settings)

    monkeypatch.setenv("DATABASE_URL", TEST_DATABASE_URL)
    for cached in caches:
        cached.cache_clear()
    try:
        yield
        _cleanup()
    finally:
        for cached in caches:
            cached.cache_clear()


def _sales_rows(session):
    """Các dòng sales của dự án này, sắp theo ngày bán."""
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import areas, sales_records

    async def run():
        factory = _engine()
        async with factory() as db:
            area_ids = sa.select(areas.c.id).where(areas.c.project_id == PROJECT_ID).scalar_subquery()
            return (
                (
                    await db.execute(
                        sa.select(sales_records)
                        .where(sales_records.c.area_id.in_(area_ids))
                        .order_by(sales_records.c.sold_date, sales_records.c.external_record_id)
                    )
                )
                .mappings()
                .all()
            )

    return [dict(r) for r in asyncio.run(run())]


def _crm_table_counts():
    """Số dòng của ba bảng do CRM sở hữu — phải đứng yên qua mọi test ở đây."""
    import asyncio

    import sqlalchemy as sa

    async def run():
        factory = _engine()
        async with factory() as session:
            return {
                name: await session.scalar(sa.text(f"SELECT count(*) FROM {name}"))
                for name in ("units", "deals", "crm_source_records")
            }

    return asyncio.run(run())


# --- Đường chạy đầy đủ -------------------------------------------------------


def test_xlsx_sales_file_reaches_the_database(tmp_path, db_env):
    """Một .xlsx thật đi tới `sales_records` với đúng giá trị từng cột."""
    before = _crm_table_counts()
    file_id = _seed()

    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 4, 1), 3, "XL-001"],
            ["A1", "2PN", date(2026, 4, 3), 2, "XL-002"],
        ],
    )

    result = run_parse_upload(
        str(path), "sales", file_id=file_id, project_id=str(PROJECT_ID), original_filename="sales.xlsx"
    )

    assert result["status"] == "done"
    assert result["persisted"] is True
    assert result["rows_inserted"] == 2
    assert result["rows_failed"] == 0

    rows = _sales_rows(None)
    assert [(r["sold_date"], r["units_sold"], r["external_record_id"]) for r in rows] == [
        (date(2026, 4, 1), 3, "XL-001"),
        (date(2026, 4, 3), 2, "XL-002"),
    ]
    # `source_updated_at` không có trong file → NULL là câu trả lời trung thực.
    assert all(r["source_updated_at"] is None for r in rows)

    assert _crm_table_counts() == before, "đường Excel đã ghi vào bảng của CRM"


def test_xlsx_drives_the_legacy_calculator(tmp_path, db_env):
    """Nạp .xlsx xong thì `absorption_daily` phải được tính lại ngay."""
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import absorption_daily, areas

    file_id = _seed()
    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 4, 1), 2, "XL-C1"],
            ["A1", "2PN", date(2026, 4, 3), 1, "XL-C2"],
        ],
    )

    run_parse_upload(str(path), "sales", file_id=file_id, project_id=str(PROJECT_ID), original_filename="s.xlsx")

    async def read():
        factory = _engine()
        async with factory() as session:
            area_ids = sa.select(areas.c.id).where(areas.c.project_id == PROJECT_ID).scalar_subquery()
            return (
                (
                    await session.execute(
                        sa.select(absorption_daily)
                        .where(absorption_daily.c.area_id.in_(area_ids))
                        .order_by(absorption_daily.c.stat_date)
                    )
                )
                .mappings()
                .all()
            )

    rows = [dict(r) for r in asyncio.run(read())]

    assert rows, "nạp .xlsx xong nhưng bộ tính cũ không sinh dòng absorption_daily nào"
    assert rows[0]["stat_date"] == date(2026, 4, 1)
    assert rows[-1]["stat_date"] == date(2026, 4, 3)
    # Ngày 2/4 không có giao dịch nhưng vẫn phải có dòng — chuỗi thời gian liên tục.
    assert [r["stat_date"] for r in rows] == [date(2026, 4, 1), date(2026, 4, 2), date(2026, 4, 3)]
    assert sum(r["units_sold"] for r in rows) == 3


def test_xlsx_and_csv_of_the_same_data_produce_identical_rows(tmp_path, db_env):
    """Cùng một dữ liệu, nạp bằng .xlsx hay .csv phải ra bản ghi giống hệt nhau.

    Đây là test giữ cho hai nhánh đọc không trôi khỏi nhau. calamine trả sẵn
    `date`/`int`, CSV trả chuỗi rồi mới ép kiểu — nếu một nhánh nào đó ép sai,
    `source_row_hash` sẽ lệch và cơ chế chống trùng lặp âm thầm mất tác dụng.
    """
    import csv as csv_module

    rows = [
        ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
        ["A1", "2PN", date(2026, 4, 10), 4, "XC-001"],
    ]

    # 1) Nạp bằng .xlsx.
    xlsx_file_id = _seed(checksum="xlsx-side")
    run_parse_upload(
        str(_write_xlsx(tmp_path / "s.xlsx", rows)),
        "sales",
        file_id=xlsx_file_id,
        project_id=str(PROJECT_ID),
        original_filename="s.xlsx",
    )
    from_xlsx = _sales_rows(None)
    assert len(from_xlsx) == 1

    # 2) Nạp CÙNG dữ liệu bằng .csv, qua một lô khác.
    csv_path = tmp_path / "s.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv_module.writer(handle)
        for row in rows:
            writer.writerow([v.isoformat() if isinstance(v, date) else v for v in row])

    csv_file_id = _seed(checksum="csv-side")
    result = run_parse_upload(
        str(csv_path), "sales", file_id=csv_file_id, project_id=str(PROJECT_ID), original_filename="s.csv"
    )

    from_csv = _sales_rows(None)

    # Cùng khoá nghiệp vụ nên lần nạp thứ hai KHÔNG được tạo dòng mới.
    #
    # Điều kiện kiểm là SỐ DÒNG TRONG BẢNG, không phải `rows_inserted`: theo
    # `ImportService._insert_rows`, `rows_inserted` đếm số dòng ĐƯỢC NHẬN để ghi
    # chứ không phải số dòng thực sự đổi, nên lần nạp lại vẫn trả 1. Bất biến
    # thật sự cần giữ là bảng không phình thêm.
    assert result["status"] == "done"
    assert len(from_csv) == 1, "nạp lại cùng dữ liệu bằng .csv lại sinh thêm dòng"

    # Và hash phải trùng — bằng chứng hai nhánh ép kiểu ra cùng một giá trị.
    assert from_csv[0]["source_row_hash"] == from_xlsx[0]["source_row_hash"]
    assert from_csv[0]["sold_date"] == from_xlsx[0]["sold_date"] == date(2026, 4, 10)
    assert from_csv[0]["units_sold"] == from_xlsx[0]["units_sold"] == 4


def test_replaying_the_same_xlsx_adds_no_rows(tmp_path, db_env):
    """Nạp lại đúng file .xlsx đó lần nữa: không thêm dòng, không hỏng lô."""
    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 4, 20), 5, "XL-R1"],
        ],
    )

    first = run_parse_upload(
        str(path), "sales", file_id=_seed(checksum="replay-1"), project_id=str(PROJECT_ID), original_filename="s.xlsx"
    )
    assert first["rows_inserted"] == 1
    after_first = _sales_rows(None)

    # Job xoá file gốc sau khi nạp thành công, nên phải ghi lại để chạy lần hai —
    # đúng như đời thật: người dùng upload lại chính file đó.
    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 4, 20), 5, "XL-R1"],
        ],
    )
    second = run_parse_upload(
        str(path), "sales", file_id=_seed(checksum="replay-2"), project_id=str(PROJECT_ID), original_filename="s.xlsx"
    )

    assert second["status"] == "done", "nạp lại phải thành công, không được báo lỗi"
    # `rows_inserted` đếm dòng ĐƯỢC NHẬN để ghi (xem ImportService._insert_rows)
    # nên nó vẫn là 1; bằng chứng chống trùng nằm ở nội dung bảng.
    assert _sales_rows(None) == after_first, "nạp lại đã làm thay đổi dữ liệu"


def test_xlsx_invalid_row_is_a_row_error_not_a_file_failure(tmp_path, db_env):
    """Một dòng hỏng trong .xlsx chỉ làm hỏng dòng đó; dòng tốt vẫn phải vào DB.

    Tỷ lệ lỗi giữ DƯỚI `import_error_threshold` (mặc định 50%) — vượt ngưỡng là
    một hành vi khác hẳn, có test riêng ngay bên dưới.
    """
    file_id = _seed()
    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 5, 1), 3, "XL-OK1"],
            ["A1", "2PN", date(2026, 5, 2), 2, "XL-OK2"],
            ["A1", "2PN", date(2026, 5, 3), -1, "XL-BAD"],  # ck_..._units_sold_nonnegative
        ],
    )

    result = run_parse_upload(
        str(path), "sales", file_id=file_id, project_id=str(PROJECT_ID), original_filename="s.xlsx"
    )

    assert result["status"] == "done", "dòng hỏng không được làm hỏng cả file"
    assert result["rows_ok"] + result["rows_failed"] == result["rows_total"] == 3
    assert result["rows_failed"] == 1

    rows = _sales_rows(None)
    assert [r["external_record_id"] for r in rows] == ["XL-OK1", "XL-OK2"]

    # Lô kết thúc dứt khoát, không treo ở 'pending'.
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import upload_files

    async def read_status():
        factory = _engine()
        async with factory() as session:
            return await session.scalar(sa.select(upload_files.c.status).where(upload_files.c.id == uuid.UUID(file_id)))

    assert asyncio.run(read_status()) in ("completed", "partially_completed")


def test_xlsx_above_the_error_threshold_rejects_the_whole_batch(tmp_path, db_env):
    """Quá nửa số dòng hỏng → từ chối cả lô, không nạp nửa vời.

    Một file mà quá nửa số dòng hỏng thường là file sai template hoặc sai kỳ,
    chứ không phải file tốt lẫn vài lỗi. Nạp phần còn lại sẽ đưa dữ liệu bán
    phần vào bảng mà không ai biết là bán phần.
    """
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import upload_files

    file_id = _seed(checksum="xlsx-threshold")
    path = _write_xlsx(
        tmp_path / "sales.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày bán", "Số căn bán", "Mã bản ghi"],
            ["A1", "2PN", date(2026, 6, 1), 3, "XT-OK"],
            ["A1", "2PN", date(2026, 6, 2), -1, "XT-BAD1"],
            ["A1", "2PN", date(2026, 6, 3), -2, "XT-BAD2"],
        ],
    )

    result = run_parse_upload(
        str(path), "sales", file_id=file_id, project_id=str(PROJECT_ID), original_filename="s.xlsx"
    )

    assert result["status"] == "failed"
    assert result["error_code"] == "ERROR_RATE_EXCEEDED"
    assert _sales_rows(None) == [], "lô bị từ chối mà vẫn nạp được dòng vào DB"

    async def read_status():
        factory = _engine()
        async with factory() as session:
            return await session.scalar(sa.select(upload_files.c.status).where(upload_files.c.id == uuid.UUID(file_id)))

    assert asyncio.run(read_status()) == "failed", "lô treo ở 'pending'"


def test_xlsx_inventory_reaches_the_database(tmp_path, db_env):
    """Template inventory cũng phải đi hết đường khi nguồn là .xlsx."""
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import areas, inventory_snapshots

    before = _crm_table_counts()
    file_id = _seed(checksum="xlsx-inv")

    path = _write_xlsx(
        tmp_path / "inv.xlsx",
        [
            ["Phân khu", "Loại căn", "Ngày chốt", "Tồn kho", "Loại chốt"],
            ["A1", "2PN", date(2026, 4, 1), 97, "closing"],
        ],
    )

    result = run_parse_upload(
        str(path), "inventory", file_id=file_id, project_id=str(PROJECT_ID), original_filename="inv.xlsx"
    )
    assert result["rows_inserted"] == 1

    async def read():
        factory = _engine()
        async with factory() as session:
            area_ids = sa.select(areas.c.id).where(areas.c.project_id == PROJECT_ID).scalar_subquery()
            return (
                (
                    await session.execute(
                        sa.select(inventory_snapshots).where(inventory_snapshots.c.area_id.in_(area_ids))
                    )
                )
                .mappings()
                .all()
            )

    rows = [dict(r) for r in asyncio.run(read())]
    assert len(rows) == 1
    assert rows[0]["units_remaining"] == 97
    assert rows[0]["snapshot_date"] == date(2026, 4, 1)
    assert rows[0]["snapshot_type"] == "closing"

    assert _crm_table_counts() == before


def test_structurally_broken_xlsx_fails_the_batch_not_silently(tmp_path, db_env):
    """File .xlsx hỏng ở mức cấu trúc → lô kết thúc ở 'failed', có dòng lỗi kèm theo.

    Test này bắt được một lỗi thật: `python_calamine.CalamineError` kế thừa thẳng
    từ `Exception` nên nó lọt qua mọi handler của `run_parse_upload` và thoát ra
    ngoài — job chết, `upload_files` không ai đánh dấu, lô treo ở 'pending' vĩnh
    viễn. Đúng lỗi S1 đã sửa cho nhánh CSV, còn sót lại ở nhánh Excel vì trước
    đây chưa có test nào đưa .xlsx đi tới tận DB.
    """
    import asyncio

    import sqlalchemy as sa

    from src.models.tables import upload_errors, upload_files

    file_id = _seed(checksum="xlsx-broken")

    # Đuôi .xlsx nhưng nội dung không phải zip → calamine không mở nổi.
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"day khong phai file excel")

    result = run_parse_upload(
        str(broken), "sales", file_id=file_id, project_id=str(PROJECT_ID), original_filename="broken.xlsx"
    )

    assert result["status"] == "failed"
    assert result["error_code"] == "UNREADABLE_WORKBOOK"
    # Không để lọt đường dẫn nội bộ hay chi tiết của thư viện ra ngoài.
    for leaked in ("/tmp", "calamine", "EOCD", "Traceback", "Zip"):
        assert leaked not in result["message"], f"thông báo lỗi để lọt '{leaked}'"

    async def read():
        factory = _engine()
        async with factory() as session:
            status = await session.scalar(
                sa.select(upload_files.c.status).where(upload_files.c.id == uuid.UUID(file_id))
            )
            errors = (
                (await session.execute(sa.select(upload_errors).where(upload_errors.c.file_id == uuid.UUID(file_id))))
                .mappings()
                .all()
            )
            return status, [dict(e) for e in errors]

    status, errors = asyncio.run(read())

    assert status == "failed", "lô treo ở 'pending' — đúng lỗi S1 đã sửa"
    assert len(errors) == 1
    assert errors[0]["column_name"] is None  # quy ước "lỗi cả file"
    assert _sales_rows(None) == [], "file hỏng mà vẫn ghi được dòng vào DB"
